"""Construye la fuente web de la cartera desde el Excel y la prelación autorizada."""

from __future__ import annotations

import argparse
import json
import math
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


GCR_NAMES = {
    "BAJA CALIFORNIA": "B. California",
    "BAJA CALIFORNIA SUR": "BC Sur",
    "CENTRAL": "Central",
    "MULEGE": "Mulegé",
    "NORESTE": "Noreste",
    "NOROESTE": "Noroeste",
    "NORTE": "Norte",
    "OCCIDENTE": "Occidental",
    "OCCIDENTAL": "Occidental",
    "ORIENTAL": "Oriental",
    "PENINSULAR": "Peninsular",
}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def ascii_key(value: object) -> str:
    normalized = unicodedata.normalize("NFD", text(value).upper())
    return "".join(char for char in normalized if unicodedata.category(char) != "Mn")


def iso_date(value: str) -> str:
    months = {
        "enero": 1, "febrero": 2, "marzo": 3, "abril": 4,
        "mayo": 5, "junio": 6, "julio": 7, "agosto": 8,
        "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12,
    }
    parts = text(value).lower().replace(" de ", " ").split()
    if len(parts) == 3 and parts[1] in months:
        return datetime(int(parts[2]), months[parts[1]], int(parts[0])).date().isoformat()
    return text(value)


def build(excel_path: Path, ranking_path: Path) -> dict:
    ranking = json.loads(ranking_path.read_text(encoding="utf-8-sig"))
    order = ranking.get("valores", {})
    decisions = ranking.get("estados", {})
    excluded = ranking.get("excluidos", {})

    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    sheet = workbook["Resumen integrado"]
    projects = []

    for row_number in range(4, sheet.max_row + 1):
        values = [sheet.cell(row_number, column).value for column in range(1, 13)]
        folio = text(values[1])
        if not folio:
            continue

        rank = int(order.get(folio, len(order) + row_number))
        raw_decision = ascii_key(decisions.get(folio, ""))
        decision = "va" if raw_decision in {"SIGUE", "VA", "CONTINUA"} else "pendiente"
        if raw_decision in {"NO", "NO VA"} or excluded.get(folio):
            decision = "no-va"

        origin = text(values[11])
        project_type = "Estratégico" if ascii_key(origin).startswith("ESTRATEG") else "Particular 2"
        gcr_key = ascii_key(values[2])

        projects.append({
            "folio": folio,
            "name": text(values[5]),
            "type": project_type,
            "region": GCR_NAMES.get(gcr_key, text(values[2]).title()),
            "state": text(values[3]).title(),
            "technology": text(values[4]).title(),
            "company": text(values[6]),
            "interestGroup": text(values[7]),
            "substation": text(values[8]),
            "interconnectionPoint": text(values[9]),
            "mw": float(values[10] or 0),
            "rank": rank,
            "priority": min(4, max(1, math.ceil(rank / 16))),
            "prioritySource": "Tramo inicial derivado de la prelación; editable por la gerencia",
            "decision": decision,
            "source": origin,
            "sourceRow": row_number,
        })

    projects.sort(key=lambda project: project["rank"])
    notes = []
    for folio, entries in ranking.get("notas", {}).items():
        for entry in entries:
            notes.append({
                "folio": folio,
                "session": text(entry.get("s")),
                "date": iso_date(text(entry.get("d"))),
                "text": text(entry.get("t")),
            })

    return {
        "sourceVersion": "excel-2026-08-04_preLacion-2026-08-05_v1",
        "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source": {
            "workbook": excel_path.name,
            "sheet": sheet.title,
            "ranking": ranking_path.name,
            "rows": len(projects),
        },
        "projects": projects,
        "notes": notes,
        "session": {
            "name": text(ranking.get("sesion")),
            "date": iso_date(text(ranking.get("fecha"))),
        },
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("excel", type=Path)
    parser.add_argument("ranking", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    payload = build(args.excel, args.ranking)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "output": str(args.output),
        "projects": len(payload["projects"]),
        "notes": len(payload["notes"]),
        "mw": sum(project["mw"] for project in payload["projects"]),
    }, ensure_ascii=False))


if __name__ == "__main__":
    main()
