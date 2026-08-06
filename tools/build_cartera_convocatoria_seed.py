"""Construye la fotografía JSON de cartera desde el resumen integrado del Excel."""

from __future__ import annotations

import argparse
import json
import unicodedata
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


EXCLUDED = {
    "VUPE-C2-0599-2026",
    "VUPE-C2-0555-2026",
    "CFE-CM2-0037-2026",
    "VUPE-C2-0679-2026",
    "VUPE-C2-0607-2026",
}

IN_ANALYSIS = {
    "CFE-CM2-0128-2026",
    "VUPE-C2-0386-2026",
    "CFE-CM2-0119-2026",
    "CFE-CM2-0152-2026",
    "CFE-CM2-0040-2026",
}

FACTIBLE = {
    "CFE-CM2-0047-2026",
    "CFE-CM2-0073-2026",
    "CFE-CM2-0129-2026",
}

DUPLICATES = {
    "CFE-CM2-0118-2026": "Barajas Solar",
    "VUPE-C2-0469-2026": "Barajas Solar",
    "CFE-CM2-0117-2026": "Malva Solar",
    "VUPE-C2-0460-2026": "Malva Solar",
    "CFE-CM2-0119-2026": "Armeria Solar",
    "CFE-CM2-0152-2026": "Armeria Solar",
    "CFE-CM2-0155-2026": "Parque Solar Lagos el Potrero",
    "VUPE-C2-0599-2026": "Parque Solar Lagos el Potrero",
}

REGIONS = {
    "BAJA CALIFORNIA": "B. California",
    "CENTRAL": "Central",
    "NORESTE": "Noreste",
    "NORTE": "Norte",
    "NOROESTE": "Noroeste",
    "OCCIDENTE": "Occidental",
    "OCCIDENTAL": "Occidental",
    "ORIENTAL": "Oriental",
    "PENINSULAR": "Peninsular",
}


def clean(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def title_case(value: str | None) -> str:
    if not value:
        return ""
    return " ".join(part.capitalize() for part in value.split())


def classification(folio: str, viability) -> str:
    if folio in IN_ANALYSIS:
        return "En análisis"
    if folio in EXCLUDED:
        return "Excluyente preseleccionado"
    if folio in FACTIBLE:
        return "Factible"
    normalized = "" if viability is None else str(viability).strip().lower()
    if normalized in {"0", "0.0"}:
        return "No viable"
    if normalized in {"1", "1.0"}:
        return "Factible"
    return "Por analizar"


def viability_label(value) -> str:
    normalized = "" if value is None else str(value).strip().lower()
    if normalized in {"0", "0.0"}:
        return "No viable"
    if normalized in {"1", "1.0"}:
        return "Viable"
    if normalized == "pendiente":
        return "Pendiente"
    return "Por analizar"


def json_value(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("previous_seed", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()

    previous = json.loads(args.previous_seed.read_text(encoding="utf-8-sig"))
    previous_projects = {item["folio"]: item for item in previous.get("projects", [])}

    workbook = load_workbook(args.workbook, data_only=True, read_only=True)
    sheet = workbook["Resumen integrado"]
    headers = [sheet.cell(3, column).value for column in range(1, 16)]
    projects = []

    for row_number in range(4, sheet.max_row + 1):
        values = [clean(sheet.cell(row_number, column).value) for column in range(1, 16)]
        row = dict(zip(headers, values))
        folio = row.get("Folio")
        if not folio:
            continue

        old = previous_projects.get(folio, {})
        analysis = row.get("Comentarios mesa técnica")
        if analysis in (0, "0"):
            analysis = None
        if folio == "CFE-CM2-0035-2026" and analysis:
            analysis = f"{analysis}. Oposición social directa."

        source_type = row.get("Origen") or ""
        project = {
            "folio": folio,
            "name": row.get("Proyecto") or folio,
            "type": "Estratégico" if source_type == "Estratégicos" else "Particular 2",
            "region": REGIONS.get(str(row.get("Región") or "").upper(), title_case(row.get("Región"))),
            "state": title_case(row.get("Entidad Federativa")),
            "technology": title_case(row.get("Tecnología")),
            "company": row.get("Razón Social"),
            "interestGroup": row.get("GEI"),
            "substation": row.get("SE de interconexión"),
            "interconnectionPoint": row.get("Punto de interconexión"),
            "mw": float(row.get("Generación neta (MW)") or 0),
            "rank": 0,
            "priority": 4,
            "prioritySource": "Prelación por capacidad neta; prioridad operativa editable por la gerencia",
            "decision": old.get("decision", "revision"),
            "analysisClassification": classification(folio, row.get("Viabilidad técnica")),
            "technicalViability": (
                "Pendiente" if folio in IN_ANALYSIS
                else "Viable" if folio in FACTIBLE
                else viability_label(row.get("Viabilidad técnica"))
            ),
            "technicalAnalysis": analysis,
            "projectSignedAt": json_value(row.get("Fecha de firma proyecto")),
            "duplicateGroup": DUPLICATES.get(folio),
            "source": source_type,
            "sourceRow": row_number,
            "_oldRank": old.get("rank", 9999),
        }
        projects.append(project)

    projects.sort(key=lambda item: (item["mw"], item["_oldRank"], item["folio"]))
    for rank, project in enumerate(projects, start=1):
        project["rank"] = rank
        project["priority"] = min(4, ((rank - 1) // 16) + 1)
        del project["_oldRank"]

    payload = {
        "sourceVersion": "pdf-2026-08-05_resultados-mesa-v5",
        "generatedAt": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source": {
            "workbook": args.workbook.name,
            "sheet": "Resumen integrado",
            "presentation": "ULtimo cmabios.pdf",
            "analysisMeeting": "Resultados de Mesa Técnica al 5 de agosto de 2026",
            "rows": len(projects),
            "method": "Prelación por capacidad; clasificación conforme al análisis SENER-CFE-CENACE",
        },
        "projects": projects,
        "notes": previous.get("notes", []),
        "session": previous.get("session", {"name": "Sesión del 4 de agosto de 2026", "date": "2026-08-04"}),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    counts = {}
    for project in projects:
        key = project["analysisClassification"]
        counts[key] = counts.get(key, 0) + 1
    print(json.dumps({
        "projects": len(projects),
        "mw": round(sum(project["mw"] for project in projects), 3),
        "strategic": sum(project["type"] == "Estratégico" for project in projects),
        "private": sum(project["type"] == "Particular 2" for project in projects),
        "analysis": counts,
        "output": str(args.output),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
