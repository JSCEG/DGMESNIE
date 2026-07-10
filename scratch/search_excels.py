import os
import openpyxl

doc_dir = r"c:\Proyectos\1.-DGMESNIE\Documentacion"
search_terms = ["huamantla", "340026100031426", "tlaxcala", "recurso"]

for filename in os.listdir(doc_dir):
    if filename.endswith(".xlsx"):
        file_path = os.path.join(doc_dir, filename)
        print(f"\nSearching in {filename}...")
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                for r_idx, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    for c_idx, val in enumerate(row, start=1):
                        if val is not None:
                            val_str = str(val).lower()
                            for term in search_terms:
                                if term in val_str:
                                    print(f"[{sheet_name}] Row {r_idx}, Col {c_idx} (Term: '{term}'): {val}")
        except Exception as e:
            print(f"Error reading {filename}: {e}")
