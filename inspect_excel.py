import openpyxl

wb_path = r"c:\Proyectos\1.-DGMESNIE\wwwroot\documents\BD_Privados_FirmesCGPT.xlsx"
wb = openpyxl.load_workbook(wb_path, read_only=True)

print("Sheets in workbook:", wb.sheetnames)

for sheet_name in ["34Py", "Tramites_Ambientales"]:
    if sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        # Get first row headers
        headers = []
        for cell in next(sheet.iter_rows(max_row=1, values_only=True)):
            if cell is not None:
                headers.append(str(cell).strip())
        print(f"\n--- Headers for sheet: {sheet_name} ({len(headers)} columns) ---")
        print(", ".join(headers))
    else:
        print(f"\nSheet {sheet_name} not found.")
